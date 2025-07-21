"""
AI-powered resume parsing service for extracting text and candidate information from uploaded files.

This service handles file upload processing, text extraction from PDF/DOCX files,
and AI-powered candidate information parsing for the RecruiterRadar system.
"""

import re
import uuid
import logging
import json
from typing import Optional, List, Dict, Any
from pathlib import Path
from datetime import datetime
import fitz  # PyMuPDF
from docx import Document
from fastapi import UploadFile, HTTPException

# OCR fallback imports
try:
    import pytesseract
    from pdf2image import convert_from_bytes

    OCR_AVAILABLE = True
except ImportError:
    OCR_AVAILABLE = False
    logger = logging.getLogger(__name__)
    logger.warning(
        "OCR libraries (pytesseract, pdf2image) not available. Image-based PDFs may fail."
    )

from ..models.candidate import CandidateProfile
from ..services.llm_service import LLMService
from ..services.ai_extraction_service import AIExtractionService
from ..models.extraction_models import ExtractedResumeData

logger = logging.getLogger(__name__)


class ResumeParsingError(Exception):
    """Custom exception for resume parsing errors."""

    pass


class ResumeParser:
    """AI-powered resume parser using GPT-4."""

    SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}
    MAX_FILE_SIZE = 10 * 1024 * 1024  # 10MB

    def __init__(self, llm_service: LLMService):
        """Initialize the AI-powered resume parser."""
        self.llm_service = llm_service
        self.ai_extractor = AIExtractionService(llm_service)

    async def parse_uploaded_file(
        self, file: UploadFile, candidate_name: Optional[str] = None
    ) -> CandidateProfile:
        """
        Parse an uploaded resume file using AI extraction.

        Args:
            file: The uploaded file object
            candidate_name: Optional name override (if not extracted from resume)

        Returns:
            CandidateProfile object ready for ingestion

        Raises:
            ResumeParsingError: If parsing fails
            HTTPException: If file validation fails
        """
        # Validate file
        await self._validate_file(file)

        # Extract text based on file type
        file_extension = Path(file.filename or "").suffix.lower()

        try:
            text = await self._extract_text_from_file(file, file_extension)

            # Use AI extraction service
            extracted_data = await self.ai_extractor.extract_resume_data(
                resume_text=text,
                timeout_seconds=45,  # Increased to 45s to prevent timeouts
            )

            if not extracted_data:
                raise ResumeParsingError("AI extraction failed")

            # Override name if provided
            if candidate_name:
                extracted_data.name = candidate_name

            # Create candidate profile from AI extracted data
            candidate_profile = CandidateProfile(
                id=f"uploaded_{uuid.uuid4().hex[:8]}",
                name=extracted_data.name,
                email=extracted_data.email,
                raw_resume_text=text,
                skills=extracted_data.technical_skills,
                experience_years=int(extracted_data.total_experience_years),
                visa_status=None,  # Not extracted by AI yet
                location=extracted_data.location,
                github_url=extracted_data.github_url,
                linkedin_url=extracted_data.linkedin_url,
            )

            logger.info(
                f"Successfully parsed resume for {candidate_profile.name} "
                f"(email: {candidate_profile.email or 'N/A'}) "
                f"(AI confidence: {extracted_data.extraction_confidence:.2f})"
            )
            return candidate_profile

        except Exception as e:
            logger.error(f"Failed to parse resume {file.filename}: {e}")
            raise ResumeParsingError(f"Resume parsing failed: {str(e)}")

    async def parse_resume(
        self, file_content: bytes, file_type: str
    ) -> Optional[ExtractedResumeData]:
        """
        Parse resume using AI extraction.

        Returns comprehensive extracted data or None if parsing fails.
        """
        # Extract text from file
        text = self.extract_text_from_file(file_content, file_type)
        if not text:
            logger.error(f"Failed to extract text from {file_type} file")
            return None

        # Use AI extraction service
        extracted_data = await self.ai_extractor.extract_resume_data(
            resume_text=text,
            timeout_seconds=45,  # Increased to 45s to prevent timeouts
        )

        return extracted_data

    async def _extract_text_from_file(
        self, file: UploadFile, file_extension: str
    ) -> str:
        """Extract text from file based on extension."""
        if file_extension == ".pdf":
            return await self._extract_text_from_pdf(file)
        elif file_extension == ".docx":
            return await self._extract_text_from_docx(file)
        elif file_extension == ".txt":
            return await self._extract_text_from_txt(file)
        else:
            raise ResumeParsingError(f"Unsupported file type: {file_extension}")

    @staticmethod
    def extract_text_from_file(file_content: bytes, file_type: str) -> Optional[str]:
        """Extract text from various file formats (static method for direct use)."""
        try:
            if file_type.lower() == "pdf":
                pdf_document = fitz.open(stream=file_content, filetype="pdf")
                text = ""
                for page_num in range(pdf_document.page_count):
                    page = pdf_document[page_num]
                    text += page.get_text() + "\n"
                pdf_document.close()
                return text.strip() if text.strip() else None

            elif file_type.lower() == "docx":
                from io import BytesIO

                doc = Document(BytesIO(file_content))
                text = ""
                for paragraph in doc.paragraphs:
                    text += paragraph.text + "\n"
                return text.strip() if text.strip() else None

            elif file_type.lower() == "txt":
                text = file_content.decode("utf-8")
                return text.strip() if text.strip() else None

            else:
                logger.error(f"Unsupported file type: {file_type}")
                return None

        except Exception as e:
            logger.error(f"Text extraction failed: {str(e)}")
            return None

    async def _validate_file(self, file: UploadFile) -> None:
        """Validate uploaded file."""
        if not file.filename:
            raise HTTPException(status_code=400, detail="No filename provided")

        file_extension = Path(file.filename).suffix.lower()
        if file_extension not in self.SUPPORTED_EXTENSIONS:
            raise HTTPException(
                status_code=400,
                detail=f"Unsupported file type. Supported: {', '.join(self.SUPPORTED_EXTENSIONS)}",
            )

        # Check file size
        content = await file.read()
        if len(content) > self.MAX_FILE_SIZE:
            raise HTTPException(status_code=400, detail="File too large (max 10MB)")

        # Reset file pointer
        await file.seek(0)

    async def _extract_text_from_pdf(self, file: UploadFile) -> str:
        """Extract text from PDF file using PyMuPDF with OCR fallback."""
        content = await file.read()

        try:
            pdf_document = fitz.open(stream=content, filetype="pdf")
            text = ""

            for page_num in range(pdf_document.page_count):
                page = pdf_document[page_num]
                page_text = page.get_text()
                text += page_text + "\n"

            pdf_document.close()

            # 🔍 CRITICAL LOGGING: See what we actually extracted
            text_length = len(text.strip())
            text_preview = (
                text.strip()[:200] + "..." if len(text.strip()) > 200 else text.strip()
            )

            logger.info(f"📄 PDF text extraction: {text_length} characters extracted")
            logger.info(f"📝 Text preview (first 200 chars): {text_preview}")

            # 🚨 OCR FALLBACK: If PyMuPDF failed (image-based PDF)
            if text_length < 100:  # Suspiciously short for a resume
                logger.warning(
                    f"⚠️ Short text extraction ({text_length} chars) - attempting OCR fallback"
                )

                if OCR_AVAILABLE:
                    try:
                        # Convert PDF to images and run OCR
                        logger.info("🔍 Converting PDF to images for OCR...")
                        images = convert_from_bytes(
                            content, first_page=1, last_page=3
                        )  # Limit to first 3 pages

                        ocr_text = ""
                        for i, image in enumerate(images):
                            logger.info(f"🔍 Running OCR on page {i+1}...")
                            page_text = pytesseract.image_to_string(image, lang="eng")
                            ocr_text += page_text + "\n"

                        ocr_length = len(ocr_text.strip())
                        if ocr_length > text_length:
                            logger.info(
                                f"✅ OCR success: {ocr_length} chars (vs {text_length} from PyMuPDF)"
                            )
                            text = ocr_text
                            text_preview = (
                                text.strip()[:200] + "..."
                                if len(text.strip()) > 200
                                else text.strip()
                            )
                            logger.info(f"📝 OCR text preview: {text_preview}")
                        else:
                            logger.warning(
                                f"❌ OCR didn't improve extraction: {ocr_length} chars"
                            )

                    except Exception as ocr_error:
                        logger.error(f"❌ OCR fallback failed: {ocr_error}")
                else:
                    logger.warning(
                        "❌ OCR libraries not available for image-based PDF fallback"
                    )

            if not text.strip():
                raise ResumeParsingError(
                    "No readable text found in PDF (tried PyMuPDF + OCR)"
                )

            return text.strip()

        except Exception as e:
            logger.error(f"PDF text extraction failed: {e}")
            raise ResumeParsingError(f"Failed to extract text from PDF: {str(e)}")

    async def _extract_text_from_docx(self, file: UploadFile) -> str:
        """Extract text from DOCX file using python-docx."""
        content = await file.read()

        try:
            # Save content to temporary file-like object
            from io import BytesIO

            doc = Document(BytesIO(content))

            text = ""
            for paragraph in doc.paragraphs:
                text += paragraph.text + "\n"

            if not text.strip():
                raise ResumeParsingError("No text could be extracted from DOCX")

            return text.strip()

        except Exception as e:
            raise ResumeParsingError(f"DOCX text extraction failed: {str(e)}")

    async def _extract_text_from_txt(self, file: UploadFile) -> str:
        """Extract text from TXT file."""
        try:
            content = await file.read()
            text = content.decode("utf-8")

            if not text.strip():
                raise ResumeParsingError("Text file is empty")

            return text.strip()

        except UnicodeDecodeError:
            raise ResumeParsingError("Could not decode text file (invalid UTF-8)")
        except Exception as e:
            raise ResumeParsingError(f"Text extraction failed: {str(e)}")
